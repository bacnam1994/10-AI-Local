"""Test 13: DOCX document generation with cover, tables, headers, and styles."""

import pytest
from docx import Document
from app.core.models import Job, SourceReference, OutputFormat, JobPlan
from app.document.docx_builder import DocxBuilder


def test_docx_builder_creates_valid_document(temp_output_dir):
    builder = DocxBuilder()
    job = Job(
        job_id="docx_test_job",
        prompt="SOP DOCX test",
        output_dir=str(temp_output_dir),
        output_formats=[OutputFormat.DOCX],
        plan=JobPlan(
            objective="QUY TRÌNH THAO TÁC CHUẨN VỆ SINH PHÒNG SẠCH CẤP D",
            tasks=[]
        )
    )

    sections = [
        ("1. Mục đích và Phạm vi", "Quy trình này quy định các bước vệ sinh cho khu vực phòng sạch."),
        ("2. Thiết bị và Hóa chất", "Sử dụng cồn 70% đã tiệt trùng qua màng lọc 0.2 micromet.")
    ]
    references = [
        SourceReference(
            reference_id="ref1",
            job_id=job.job_id,
            title="Thông tư 35/2018/TT-BYT",
            status="Có hiệu lực"
        )
    ]

    out_file = temp_output_dir / "test_sop.docx"
    builder.build(
        job=job,
        sections=sections,
        references=references,
        output_file=out_file,
        doc_code="SOP-TEST-001",
        version="1.0"
    )

    assert out_file.exists()
    assert out_file.stat().st_size > 5000  # Non-trivial DOCX file

    # Read back generated DOCX using python-docx to verify integrity
    doc = Document(str(out_file))
    paragraphs_text = [p.text for p in doc.paragraphs if p.text]
    full_text = " ".join(paragraphs_text)

    assert "QUY TRÌNH THAO TÁC CHUẨN" in full_text
    assert "1. Mục đích và Phạm vi" in full_text
    assert len(doc.tables) >= 3  # Metadata, Revision History, References tables

    # Check footer text
    footer_text = doc.sections[0].footer.paragraphs[0].text
    assert "DRAFT — CHỜ RÀ SOÁT VÀ PHÊ DUYỆT" in footer_text
