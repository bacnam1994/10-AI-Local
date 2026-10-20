"""DOCX Document Builder using python-docx."""

import re
from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.core.models import Job, SourceReference


class DocxBuilder:
    """Constructs professional Word (.docx) documents with cover, TOC, tables, and styles."""

    def __init__(self):
        pass

    def build(
        self,
        job: Job,
        sections: list[tuple[str, str]],  # List of (title, markdown_content)
        references: list[SourceReference],
        output_file: Path,
        doc_code: str = "SOP-PROD-001",
        version: str = "1.0",
        effective_date: str = "2026-10-20",
        department: str = "Phòng Đảm bảo Chất lượng (QA)",
    ) -> Path:
        """Assemble full DOCX document."""
        doc = Document()

        # Set page margins (1 inch = 72pt)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # 1. Setup Header & Footer
        self._setup_header_footer(doc, doc_code, version)

        # 2. Cover / Title Block
        self._add_cover_block(doc, job, doc_code, version, effective_date, department)

        # 3. Revision History Table
        self._add_revision_history_table(doc, version, effective_date)

        # 4. Regulatory References Table
        if references:
            self._add_references_table(doc, references)

        # 5. Table of Contents Placeholder Block
        self._add_toc_block(doc)

        # Page break before main body
        doc.add_page_break()

        # 6. Main Body Sections
        for title, md_content in sections:
            self._render_section(doc, title, md_content)

        output_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_file))
        return output_file

    def _setup_header_footer(self, doc: Document, doc_code: str, version: str) -> None:
        """Configure running headers and footers."""
        section = doc.sections[0]

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"MÃ: {doc_code} | PHIÊN BẢN: {version} | DRAFT")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("TRẠNG THÁI: DRAFT — CHỜ RÀ SOÁT VÀ PHÊ DUYỆT (CHƯA CÓ GIÁ TRỊ PHÁP LÝ CHÍNH THỨC)")
        frun.font.size = Pt(8)
        frun.font.italic = True
        frun.font.color.rgb = RGBColor(180, 50, 50)

    def _add_cover_block(
        self,
        doc: Document,
        job: Job,
        doc_code: str,
        version: str,
        effective_date: str,
        department: str
    ) -> None:
        """Add prominent header and cover information block."""
        # Top banner
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_org = title_p.add_run("TỔNG CÔNG TY DƯỢC PHẨM VIỆT NAM\n")
        run_org.bold = True
        run_org.font.size = Pt(13)
        run_dept = title_p.add_run(f"{department.upper()}\n")
        run_dept.font.size = Pt(10)
        run_dept.font.color.rgb = RGBColor(90, 90, 90)

        # Main Title
        objective = job.plan.objective if job.plan else "QUY TRÌNH THAO TÁC CHUẨN (SOP)"
        main_title_p = doc.add_paragraph()
        main_title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        main_title_run = main_title_p.add_run(f"\n{objective.upper()}\n")
        main_title_run.bold = True
        main_title_run.font.size = Pt(18)
        main_title_run.font.color.rgb = RGBColor(24, 76, 120)

        # Draft Banner Box
        draft_p = doc.add_paragraph()
        draft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        draft_run = draft_p.add_run("— BẢN THẢO (DRAFT) — CHỜ RÀ SOÁT VÀ PHÊ DUYỆT —\n")
        draft_run.bold = True
        draft_run.font.size = Pt(10.5)
        draft_run.font.color.rgb = RGBColor(190, 40, 40)

        # Metadata Table
        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        meta_rows = [
            ("Mã tài liệu (Document Code):", doc_code),
            ("Phiên bản (Version):", f"{version} (Bản thảo sơ bộ)"),
            ("Ngày dự thảo (Date):", effective_date),
            ("Bộ phận soạn thảo:", department)
        ]

        for i, (label, val) in enumerate(meta_rows):
            row = table.rows[i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(2.5)
            c1.width = Inches(4.0)

            p0 = c0.paragraphs[0]
            r0 = p0.add_run(label)
            r0.bold = True
            r0.font.size = Pt(9.5)

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(val)
            r1.font.size = Pt(9.5)

        doc.add_paragraph()

    def _add_revision_history_table(self, doc: Document, version: str, effective_date: str) -> None:
        """Add formal Revision History table."""
        h = doc.add_paragraph()
        hrun = h.add_run("BẢNG LỊCH SỬ SỬA ĐỔI (REVISION HISTORY)")
        hrun.bold = True
        hrun.font.size = Pt(11)
        hrun.font.color.rgb = RGBColor(30, 30, 30)

        table = doc.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Phiên bản", "Ngày ban hành", "Người biên soạn", "Tóm tắt nội dung thay đổi"]
        
        # Header row
        hdr_row = table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(9)
            # Background shading
            shading = parse_xml(r'<w:shd {} w:fill="E8EEF5"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)

        # Initial draft row
        data_row = table.rows[1]
        data_vals = [version, effective_date, "AI Orchestrator (Bản thảo)", "Khởi tạo bản thảo tài liệu tự động theo yêu cầu"]
        for idx, val in enumerate(data_vals):
            cell = data_row.cells[idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)

        doc.add_paragraph()

    def _add_references_table(self, doc: Document, references: list[SourceReference]) -> None:
        """Add formal references and regulatory compliance table."""
        h = doc.add_paragraph()
        hrun = h.add_run("DANH MỤC TÀI LIỆU VÀ QUY CHUẨN THAM CHIẾU")
        hrun.bold = True
        hrun.font.size = Pt(11)
        hrun.font.color.rgb = RGBColor(30, 30, 30)

        table = doc.add_table(rows=len(references) + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["STT", "Số hiệu / Tiêu đề", "Cơ quan ban hành", "Tình trạng hiệu lực"]

        hdr_row = table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(9)
            shading = parse_xml(r'<w:shd {} w:fill="E8EEF5"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)

        for i, ref in enumerate(references):
            row = table.rows[i + 1]
            vals = [
                str(i + 1),
                f"{ref.doc_number + ': ' if ref.doc_number else ''}{ref.title}",
                ref.issuing_body or "Bộ Y Tế",
                ref.status
            ]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(8.5)

        doc.add_paragraph()

    def _add_toc_block(self, doc: Document) -> None:
        """Add Table of Contents placeholder with Word TOC field code."""
        h = doc.add_paragraph()
        hrun = h.add_run("MỤC LỤC TÀI LIỆU")
        hrun.bold = True
        hrun.font.size = Pt(11)
        hrun.font.color.rgb = RGBColor(30, 30, 30)

        p = doc.add_paragraph()
        run = p.add_run("[Nhấn F9 hoặc click chuột phải chọn 'Update Field' để cập nhật Mục Lục tự động trong MS Word]")
        run.font.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(120, 120, 120)

        # Word TOC Field Code
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        p._p.append(fldChar1)
        p._p.append(instrText)
        p._p.append(fldChar2)
        p._p.append(fldChar3)

    def _render_section(self, doc: Document, title: str, md_text: str) -> None:
        """Render a section title and parse markdown lines into paragraphs, sub-headings, and lists."""
        # Main Section Heading (Heading 1)
        h1 = doc.add_heading(level=1)
        r1 = h1.add_run(title)
        r1.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = RGBColor(24, 76, 120)

        lines = md_text.splitlines()
        in_table = False
        table_lines = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            # Check if this line is part of a markdown table
            if stripped.startswith("|") and stripped.endswith("|"):
                in_table = True
                table_lines.append(stripped)
                continue
            elif in_table:
                # Flush markdown table
                self._render_markdown_table(doc, table_lines)
                table_lines = []
                in_table = False

            # Sub-heading 2
            if stripped.startswith("### "):
                h = doc.add_heading(level=2)
                r = h.add_run(stripped[4:])
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(40, 40, 40)
            # Sub-heading 3
            elif stripped.startswith("#### "):
                h = doc.add_heading(level=3)
                r = h.add_run(stripped[5:])
                r.bold = True
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(60, 60, 60)
            # Bullet list
            elif stripped.startswith("- ") or stripped.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                self._add_formatted_text(p, stripped[2:])
            # Numbered list
            elif re.match(r"^\d+\.\s+", stripped):
                p = doc.add_paragraph(style="List Number")
                match = re.match(r"^\d+\.\s+", stripped)
                self._add_formatted_text(p, stripped[match.end():])
            # Normal paragraph
            else:
                p = doc.add_paragraph()
                self._add_formatted_text(p, stripped)

        if in_table and table_lines:
            self._render_markdown_table(doc, table_lines)

    def _render_markdown_table(self, doc: Document, lines: list[str]) -> None:
        """Parse markdown table rows and append a docx Table."""
        rows_data = []
        for line in lines:
            # Ignore separator lines like |---|---|
            if re.match(r"^\|(?:\s*[-:]+[-|\s:]*)\|$", line):
                continue
            cols = [c.strip() for c in line.strip("|").split("|")]
            rows_data.append(cols)

        if not rows_data:
            return

        col_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, r_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for c_idx, val in enumerate(r_data):
                if c_idx < col_count:
                    cell = row.cells[c_idx]
                    p = cell.paragraphs[0]
                    r = p.add_run(val)
                    r.font.size = Pt(9)
                    if is_header:
                        r.bold = True
                        shd = parse_xml(r'<w:shd %s w:fill="F0F4F8"/>' % nsdecls('w'))
                        cell._tc.get_or_add_tcPr().append(shd)

        doc.add_paragraph()

    def _add_formatted_text(self, paragraph, text: str) -> None:
        """Parse bold (**text**) within paragraph."""
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
                run.font.size = Pt(10)
            else:
                run = paragraph.add_run(part)
                run.font.size = Pt(10)
